from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook


OUTPUT_DIR = Path("output")
EXCEL_FILE = OUTPUT_DIR / "rag_results.xlsx"


HEADERS = [
    "Timestamp",
    "User Input",
    "AI Output",
    "Status",
    "Source Files",
    "Categories",
    "Pages",
    "Latency Seconds"
]


def initialize_excel():

    OUTPUT_DIR.mkdir(exist_ok=True)

    if not EXCEL_FILE.exists():

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "RAG Results"

        worksheet.append(HEADERS)

        workbook.save(EXCEL_FILE)


def log_rag_result(question, answer, status, documents, latency=0):
    initialize_excel()

    try:
        workbook = load_workbook(EXCEL_FILE)
        worksheet = workbook["RAG Results"]

        source_files = []
        categories = []
        pages = []

        for document in documents:
            metadata = document.metadata

            source_file = metadata.get("source_file")
            category = metadata.get("document_category")
            page = metadata.get("page_number")

            if source_file:
                source_files.append(str(source_file))

            if category:
                categories.append(str(category))

            if page is not None:
                pages.append(str(page))

        # Remove duplicates
        source_files = list(dict.fromkeys(source_files))
        categories = list(dict.fromkeys(categories))
        pages = list(dict.fromkeys(pages))

        # Convert lists to Excel-compatible strings
        source_files_text = ", ".join(source_files)
        categories_text = ", ".join(categories)
        pages_text = ", ".join(pages)

        worksheet.append([
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            str(question),
            str(answer),
            str(status),
            source_files_text,
            categories_text,
            pages_text,
            float(latency)
        ])

        workbook.save(EXCEL_FILE)

        return True

    except PermissionError:
        print(
            "WARNING: rag_results.xlsx is currently open or locked. "
            "Please close the Excel file and try again."
        )
        return False

    except Exception as error:
        print(f"Excel logging failed: {error}")
        return False